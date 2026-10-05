"""Generación de Excel (xlsx) con las inscripciones validadas."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Columnas: (encabezado, atributo del modelo Postulante)
COLUMNAS = [
    ("N°", "id"),
    ("Apellido", "apellido"),
    ("Nombre", "nombre"),
    ("DNI", "dni"),
    ("CUIL", "cuil"),
    ("Género", "sexo"),
    ("Fecha de nacimiento", "fecha_nacimiento"),
    ("Nacionalidad", "nacionalidad"),
    ("Tel. celular", "telefono_celular"),
    ("Tel. particular", "telefono_particular"),
    ("Tel. alternativo", "telefono_alternativo"),
    ("Email", "email"),
    ("Domicilio real - Calle", "real_calle"),
    ("Real - Número", "real_numero"),
    ("Real - Piso/Depto", "real_piso_depto"),
    ("Real - Código Postal", "real_codigo_postal"),
    ("Real - Localidad", "real_localidad"),
    ("Real - Provincia", "real_provincia"),
    ("Domicilio constituido - Calle", "const_calle"),
    ("Constituido - Número", "const_numero"),
    ("Constituido - Piso/Depto", "const_piso_depto"),
    ("Constituido - Código Postal", "const_codigo_postal"),
    ("Constituido - Localidad", "const_localidad"),
    ("Constituido - Provincia", "const_provincia"),
    ("Título", "titulo"),
    ("Universidad", "universidad"),
    ("Matrícula Profesional", "matricula_profesional"),
    ("Expedida por", "expedida_por"),
    ("Especialidad", "especialidad"),
    ("Cargo actual - Establecimiento", "cargo_establecimiento"),
    ("Cargo actual - Cargo", "cargo_cargo"),
    ("Apoderado - Nombre y Apellido", "apoderado_nombre"),
    ("Apoderado - Tipo doc.", "apoderado_tipo_documento"),
    ("Apoderado - Documento", "apoderado_documento"),
    ("Apoderado - N° Acta", "apoderado_numero_acta"),
    ("Fecha de preinscripción", "creado_en"),
    ("Validada por", "validado_por"),
    ("Fecha de validación", "validado_en"),
]

SEXO_LABELS = {"M": "Masculino", "F": "Femenino", "NS": "Prefiero no decirlo"}


def _valor(p, attr):
    v = getattr(p, attr, "")
    if v is None:
        return ""
    if attr == "sexo":
        return SEXO_LABELS.get(v, v)
    if attr == "fecha_nacimiento":
        try:
            return v.strftime("%d/%m/%Y")
        except AttributeError:
            return str(v)
    if attr in ("creado_en", "validado_en"):
        try:
            return v.strftime("%d/%m/%Y %H:%M")
        except AttributeError:
            return str(v)
    return str(v)


def generar_excel_validadas(postulantes) -> bytes:
    """Genera un xlsx con las inscripciones validadas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inscripciones validadas"

    azul = "153244"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=azul, end_color=azul, fill_type="solid")

    # Encabezados
    for col_idx, (titulo, _) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col_idx, value=titulo)
        celda.font = header_font
        celda.fill = header_fill
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Filas de datos
    for fila_idx, p in enumerate(postulantes, start=2):
        for col_idx, (_, attr) in enumerate(COLUMNAS, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=_valor(p, attr))

    # Ancho de columnas aproximado según el encabezado
    for col_idx, (titulo, _) in enumerate(COLUMNAS, start=1):
        ancho = min(max(len(titulo) + 2, 12), 36)
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    ws.freeze_panes = "A2"  # fija la fila de encabezados

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
